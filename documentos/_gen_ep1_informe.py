# -*- coding: utf-8 -*-
from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path.home() / ".cursor/skills/duoc-presentations/templates/word/duoc-informe-general.docx"
OUT = Path(__file__).resolve().parent / "EP1-ARY1102-Informe-Tecnico-FreshBox.docx"

NEGRO = RGBColor(0x1A, 0x1A, 0x1A)
TEXTO = RGBColor(0x2B, 0x2B, 0x2B)
GRIS = RGBColor(0x71, 0x71, 0x7A)
AMARILLO = "FFB800"
HEADER_BG = "1A1A1A"
ALT_ROW = "F0F4F8"

HOY = date.today().strftime("%d/%m/%Y")
ALUMNO = "Matías Fernández Z."
SECCION = "ARY1102 — Arquitectura Cloud"
EMAIL = "ma.fernandezz@duocuc.cl"


def set_run_font(run, size=11, bold=False, color=TEXTO, name="Aptos"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def shade_cell(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def cant_split_row(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def style_header_row(row):
    cant_split_row(row)
    for cell in row.cells:
        shade_cell(cell, HEADER_BG)
        set_cell_margins(cell)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                set_run_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "18")
        bottom.set(qn("w:color"), AMARILLO)
        borders.append(bottom)
        tcPr.append(borders)


def style_body_row(row, alt=False):
    cant_split_row(row)
    for cell in row.cells:
        if alt:
            shade_cell(cell, ALT_ROW)
        set_cell_margins(cell)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                set_run_font(run, size=10, color=TEXTO)


def add_corp_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        hdr.cells[i].text = h
    style_header_row(hdr)
    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        for ci, val in enumerate(row_data):
            row.cells[ci].text = str(val)
        style_body_row(row, alt=(ri % 2 == 1))
    doc.add_paragraph("")


def add_h1(doc, text):
    p = doc.add_paragraph(text, style="Heading 1")
    for run in p.runs:
        set_run_font(run, size=16, bold=True, color=NEGRO)
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), AMARILLO)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_h2(doc, text):
    p = doc.add_paragraph(text, style="Heading 2")
    for run in p.runs:
        set_run_font(run, size=13, bold=True, color=NEGRO)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)


def add_p(doc, text, bullet=False):
    p = doc.add_paragraph()
    run = p.add_run(("• " if bullet else "") + text)
    set_run_font(run, size=11, color=TEXTO)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if bullet:
        p.paragraph_format.left_indent = Cm(0.5)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cant_split_row(table.rows[0])
    shade_cell(cell, "F4F5F6")
    set_cell_margins(cell, 140, 140, 180, 180)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:color"), AMARILLO)
    borders.append(left)
    for side in ("top", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "nil")
        borders.append(b)
    tcPr.append(borders)
    cell.paragraphs[0].clear()
    r1 = cell.paragraphs[0].add_run(title)
    set_run_font(r1, size=10, bold=True, color=NEGRO)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    set_run_font(r2, size=10, color=TEXTO)
    doc.add_paragraph("")


def clear_body_after_cover(doc):
    cut = None
    for p in doc.paragraphs:
        if p.style.name == "Heading 1":
            cut = p._element
            break
    if cut is None:
        return
    el = cut
    while el is not None:
        nxt = el.getnext()
        if el.tag == qn("w:sectPr"):
            break
        parent = el.getparent()
        if parent is None:
            break
        parent.remove(el)
        el = nxt
        if el is not None and el.tag == qn("w:sectPr"):
            break


def main():
    shutil.copy(TEMPLATE, OUT)
    doc = Document(str(OUT))

    mapping_idx = {
        2: ("INFORME TÉCNICO", 11, True, GRIS),
        3: ("Evaluación Parcial N°1 — ARY1102 Arquitectura Cloud", 22, True, NEGRO),
        4: (
            f"FreshBox SpA — Catálogo Online de Productos Orgánicos\nEstudiante: {ALUMNO} · {SECCION} · {HOY}",
            11,
            False,
            GRIS,
        ),
        6: ("Duoc UC · Instituto Profesional · 2026 · #VerdaderamenteConectad@s", 10, False, GRIS),
    }
    for idx, (text, size, bold, color) in mapping_idx.items():
        p = doc.paragraphs[idx]
        p.clear()
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)

    if doc.tables:
        meta = [
            ("Código Documental:", "INF-ARY1102-EP1-2026"),
            ("Versión:", "1.0"),
            ("Asignatura:", "ARY1102 Arquitectura Cloud"),
            ("Estudiante:", f"{ALUMNO} ({EMAIL})"),
            ("Fecha:", HOY),
            ("Entorno:", "AWS Academy Learner Lab · us-east-1 · workspace clases"),
        ]
        t0 = doc.tables[0]
        for ri, (k, v) in enumerate(meta):
            if ri < len(t0.rows):
                t0.rows[ri].cells[0].text = k
                t0.rows[ri].cells[1].text = v

    # Remove leftover sample tables after meta (versions, synthesis, etc.)
    # Keep only first table (ficha)
    while len(doc.tables) > 1:
        tbl = doc.tables[1]._tbl
        tbl.getparent().remove(tbl)

    clear_body_after_cover(doc)


    add_h1(doc, "Control de Versiones y Aprobaciones")
    add_corp_table(
        doc,
        ["Versión", "Fecha", "Elaborado por", "Revisado por", "Estado"],
        [
            ["1.0", HOY, ALUMNO, "Docente ARY1102", "● Entrega EP1"],
            ["1.1", HOY, ALUMNO, "—", "● Aclara pauta lab vs TO-BE empresarial"],
        ],
    )

    add_h1(doc, "Índice de Contenidos")
    for line in [
        "Introducción",
        "1.1 Fundamentación del rol del arquitecto cloud (IE1.1.2)",
        "1.2 Pilares del Well-Architected Framework (IE1.2.1)",
        "1.3 Análisis de arquitectura según Well-Architected (IE1.2.2)",
        "1.4 Priorización de requerimientos (IE1.3.2)",
        "1.5 Comparación de modelos de nube (IE2.1.1)",
        "1.6 Justificación del modelo cloud (IE2.1.2)",
        "1.7 Validación del diseño (IE2.2.2) — implementación EP1 y TO-BE empresarial",
        "Conclusiones",
        "Glosario",
        "Referencias (APA 7)",
    ]:
        add_p(doc, line)

    add_h1(doc, "Introducción")
    add_p(
        doc,
        "FreshBox SpA comercializa productos orgánicos y saludables con despacho a domicilio en la "
        "Región Metropolitana. El crecimiento sostenido del 40% trimestral exige modernizar la "
        "plataforma para publicar y administrar un catálogo online (CRUD de productos). El carrito "
        "y el procesamiento de órdenes quedan para fases posteriores.",
    )
    add_p(
        doc,
        "Este informe corresponde al Componente 1 (encargo técnico) de la EP1 de ARY1102. El "
        "repositorio Matishac14/ARY1102_Arquitectura_Cloud contiene la implementación demostrable "
        "en AWS Academy Learner Lab: Terraform modular, ECR, cinco contenedores Docker ARM64, "
        "ASG Multi-AZ, MariaDB en EC2 y GitHub Actions semi-automático.",
    )
    add_callout(
        doc,
        "Aclaración obligatoria para la defensa (pauta vs evolución)",
        "La implementación de laboratorio cumple la pauta académica. El TO-BE empresarial elimina "
        "restricciones de Learner Lab y reduce riesgo operacional mediante servicios administrados "
        "(por ejemplo ECS/Fargate y RDS Multi-AZ). En la demo se presenta lo exigido por la pauta "
        "(EC2+Docker, ASG 2–4, MySQL en EC2, ALB, VPC /22, seis subredes, SG por capa y CRUD vía ALB). "
        "ECS/Fargate y RDS se explican como evolución recomendada, no como sustituto de la demo.",
    )

    add_h1(doc, "1.1 Fundamentación del rol del arquitecto cloud (IE1.1.2)")
    add_p(
        doc,
        "El arquitecto cloud traduce objetivos de negocio en una arquitectura segura, escalable, "
        "resiliente, operable y financieramente sostenible. No solo elige servicios AWS: levanta "
        "requerimientos, evalúa riesgos, documenta decisiones y valida valor medible. En FreshBox "
        "debe sostener el crecimiento del 40% trimestral sin degradar el catálogo ni elevar CAPEX.",
    )
    add_h2(doc, "Responsabilidades técnicas")
    for t in [
        "Diseñar VPC Multi-AZ (/22), seis subredes, IGW, NAT, ALB y Security Groups por capa.",
        "Definir cómputo App con EC2 Amazon Linux + Docker y ASG (mín. 2 / máx. 4) según pauta EP1.",
        "Ubicar MySQL/MariaDB en EC2 de capa Data privada, con EBS cifrado y AWS Backup.",
        "Publicar cinco imágenes en ECR y asegurar CRUD vía ALB (frontend + cuatro APIs).",
        "Automatizar con Terraform y despliegue semi-auto (GitHub Actions), respetando LabRole.",
    ]:
        add_p(doc, t, bullet=True)
    add_h2(doc, "Responsabilidades estratégicas y de gobierno")
    for t in [
        "Alinear decisiones al negocio (time-to-market del catálogo, continuidad y costo controlado).",
        "Separar claramente lo demostrable en el lab de la hoja de ruta productiva (TO-BE).",
        "Definir criterios de aceptación: targets healthy, Multi-AZ, CRUD end-to-end ALB→EC2→MySQL.",
        "Gobernar etiquetado, presupuestos del lab, destroy al cerrar sesión y secretos temporales.",
    ]:
        add_p(doc, t, bullet=True)

    add_h1(doc, "1.2 Pilares del Well-Architected Framework (IE1.2.1)")
    add_p(
        doc,
        "Se describen los seis pilares del AWS Well-Architected Framework y su aplicación a FreshBox, "
        "distinguiendo lo ya implementado en el laboratorio de lo propuesto para producción.",
    )
    add_corp_table(
        doc,
        ["Pilar", "En la implementación EP1 (lab)", "Evolución TO-BE empresarial"],
        [
            ["Excelencia operativa", "Terraform, Actions, validate-ep1, recycle ASG", "CI/CD con pruebas, escaneo, rollback, runbooks"],
            ["Seguridad", "SG por capa, EBS cifrado, BD solo desde App, LabRole", "WAF, HTTPS/ACM, Secrets Manager, CloudTrail, SSM"],
            ["Confiabilidad", "Multi-AZ, ALB, ASG min 2, Backup 7 días", "RDS Multi-AZ, NAT por AZ, pruebas de restauración"],
            ["Rendimiento", "t4g.small ARM, 5 contenedores, ASG hasta 4", "ECS/Fargate por servicio, autoescalado por métrica"],
            ["Costos", "1 NAT, destroy, on-demand locks", "Presupuestos, VPC endpoints, right-sizing continuo"],
            ["Sostenibilidad", "Graviton, menos overprovisioning", "Apagado no-prod, retención de logs adecuada"],
        ],
    )

    add_h1(doc, "1.3 Análisis de arquitectura según Well-Architected (IE1.2.2)")
    add_p(
        doc,
        "La arquitectura del repositorio es adecuada para demostrar los conceptos de la evaluación. "
        "Una lectura Well-Architected identifica riesgos del lab y oportunidades de madurez productiva "
        "sin invalidar el cumplimiento de la pauta.",
    )
    add_corp_table(
        doc,
        ["Pilar", "Riesgo / brecha (lab o base)", "Mejora", "Prioridad"],
        [
            ["Confiabilidad", "MySQL en un solo EC2 (SPOF de datos)", "RDS Multi-AZ + backups automáticos (TO-BE)", "Alta"],
            ["Confiabilidad", "Un solo NAT Gateway", "NAT por AZ o VPC endpoints (TO-BE)", "Media"],
            ["Rendimiento", "Cinco contenedores compiten en el mismo host", "ECS/Fargate por servicio (TO-BE)", "Alta"],
            ["Seguridad", "ALB solo HTTP (sin TLS en lab)", "HTTPS + ACM + WAF (TO-BE)", "Alta"],
            ["Excelencia operativa", "Session Token caduca en Actions", "Renovación de Secrets; luego OIDC fuera del lab", "Alta"],
            ["Costos", "Olvidar destroy deja NAT/ALB/EC2 encendidos", "Checklist de cierre + alarmas de presupuesto", "Alta"],
        ],
    )
    add_p(
        doc,
        "La mejora más importante es distinguir arquitectura didáctica (EC2+Docker+MySQL en EC2) de "
        "arquitectura operable (ECS/Fargate + RDS). Ambas coexisten en este informe: la primera se "
        "demuestra; la segunda se recomienda.",
    )

    add_h1(doc, "1.4 Priorización de requerimientos (IE1.3.2)")
    add_corp_table(
        doc,
        ["ID", "Requerimiento", "Tipo", "Prioridad", "Cómo se cubre en EP1 (lab)"],
        [
            ["RF-01", "Consultar productos (GET)", "Funcional", "Alta", "get-products detrás del ALB"],
            ["RF-02", "Crear / actualizar / eliminar", "Funcional", "Alta", "create/update/delete + frontend"],
            ["RF-03", "Cinco contenedores Docker", "Funcional", "Alta", "frontend + 4 APIs en cada EC2 App"],
            ["RNF-01", "HA Multi-AZ", "No funcional", "Alta", "6 subredes, ALB, ASG min 2"],
            ["RNF-02", "Escalabilidad automática", "No funcional", "Alta", "ASG máx. 4"],
            ["RNF-03", "Seguridad por capas + cifrado", "No funcional", "Alta", "SG ALB→App→Data; EBS encrypted"],
            ["RNF-04", "Recuperación de datos", "No funcional", "Alta", "AWS Backup 7 días sobre EC2 BD"],
            ["RNF-05", "Costos controlados (lab)", "No funcional", "Alta", "1 NAT; destroy; t4g.small"],
            ["RNF-06", "Despliegues repetibles", "No funcional", "Media", "Terraform + GitHub Actions"],
            ["RF-04", "Carrito y órdenes", "Funcional", "Baja", "Fuera de alcance EP1"],
            ["RNF-07", "HTTPS / WAF / Secrets Manager", "No funcional", "Media", "Planificado en TO-BE empresarial"],
            ["RNF-08", "Cómputo y BD administrados", "No funcional", "Media", "ECS/Fargate + RDS en TO-BE"],
        ],
    )

    add_h1(doc, "1.5 Comparación de modelos de nube (IE2.1.1)")
    add_corp_table(
        doc,
        ["Criterio", "Pública", "Privada", "Híbrida"],
        [
            ["Costos", "Baja inversión inicial; OPEX; requiere gobierno", "Alto CAPEX (HW, licencias, operación)", "Costos duplicados de integración"],
            ["Seguridad", "Controles maduros; responsabilidad compartida", "Control físico total; todo a cargo propio", "Complejidad de políticas duales"],
            ["Escalabilidad", "Elástica (ASG/ECS, Multi-AZ)", "Limitada por capacidad comprada", "Condicionada por enlace y distribución"],
            ["Ventaja FreshBox", "Ideal para MVP y crecimiento 40%", "Sin ventaja clara en etapa 1", "Solo si hay legado inamovible"],
            ["Desventaja FreshBox", "Dependencia del proveedor / límites lab", "CAPEX y lead time altos", "Mayor costo y complejidad operativa"],
        ],
    )

    add_h1(doc, "1.6 Justificación del modelo cloud (IE2.1.2)")
    add_h2(doc, "Modelo seleccionado: nube pública AWS")
    add_p(
        doc,
        "Se selecciona nube pública AWS (Learner Lab para la demo; cuenta productiva para el TO-BE). "
        "AWS ofrece VPC, ALB, EC2, ASG, ECR, Backup y, en evolución, ECS, RDS, WAF e IAM avanzado "
        "sin construir plataforma física.",
    )
    add_h2(doc, "Justificación técnica")
    for t in [
        "Cumple el TO-BE académico de tres capas con componentes nativos exigidos por la pauta.",
        "Permite contenedores ARM64, Multi-AZ y balanceo sin rediseñar el dominio CRUD.",
        "IaC (Terraform) y Actions hacen reproducible el laboratorio y preparan la madurez CI/CD.",
    ]:
        add_p(doc, t, bullet=True)
    add_h2(doc, "Justificación financiera (CAPEX → OPEX)")
    add_p(
        doc,
        "Se sustituye inversión en servidores y balanceadores (CAPEX) por consumo del lab/cloud (OPEX). "
        "Servicios administrados en el TO-BE reducen horas de parches de SO/motor DB. OPEX exige "
        "igual tagging, presupuestos y destroy: no es costo automáticamente bajo.",
    )
    add_h2(doc, "Justificación estratégica")
    add_p(
        doc,
        "FreshBox acelera el time-to-market del catálogo y deja una ruta incremental hacia carrito, "
        "órdenes y endurecimiento (HTTPS, WAF, RDS, ECS) sin abandonar la separación por capas.",
    )

    add_h1(doc, "1.7 Validación del diseño (IE2.2.2)")
    add_callout(
        doc,
        "Dos conceptos que no deben confundirse en la defensa",
        "Implementación EP1 obligatoria = diseño exigido por la pauta (EC2+Docker, ASG, MySQL en EC2, "
        "ALB, ECR, VPC /22, seis subredes, SG por capa, CRUD). TO-BE empresarial sin restricciones = "
        "evolución recomendada (ECS/Fargate, RDS Multi-AZ, WAF, Secrets Manager, CloudWatch/CloudTrail, "
        "CI/CD, NAT por AZ, VPC endpoints). Presentar solo ECS/RDS sin evidenciar el lab puede afectar "
        "los indicadores de demostración (red/HA, cinco contenedores en EC2, conectividad ALB→EC2→MySQL).",
    )

    add_h2(doc, "1.7.1 Implementación EP1 obligatoria (lo que se demuestra en Academy)")
    add_p(
        doc,
        "La solución desplegada en el repositorio y validada en el Learner Lab cumple literalmente "
        "los componentes de la pauta de presentación/demostración:",
    )
    add_corp_table(
        doc,
        ["Exigencia de la pauta", "Evidencia en la implementación", "Estado"],
        [
            ["VPC CIDR /22", "VPC 10.0.0.0/22 (Terraform module network)", "● Cumple"],
            ["Seis subredes en dos AZ", "2 public + 2 private-app + 2 private-data", "● Cumple"],
            ["IGW + NAT Gateway", "Routing module; 1 NAT (presupuesto lab)", "● Cumple"],
            ["ALB público", "ALB internet-facing :80; TG health /", "● Cumple"],
            ["EC2 Amazon Linux + Docker", "ASG t4g.small AL2023 ARM; user-data Docker", "● Cumple"],
            ["ASG min 2 / max 4", "application-autoscaling module", "● Cumple"],
            ["EC2 MySQL capa privada", "EC2 MariaDB en subnet data + Backup 7 días", "● Cumple"],
            ["Cinco contenedores Docker", "frontend + get/create/update/delete en ECR", "● Cumple"],
            ["CRUD vía ALB", "http://ALB/ y /api/products (GET/POST/PUT/DELETE)", "● Cumple"],
            ["SG por capa", "ALB:80←Internet; App:80←ALB; Data:3306←App", "● Cumple"],
            ["Flujo ALB → EC2 → MySQL", "Demo consola + smoke Actions / validate-ep1", "● Cumple"],
        ],
    )
    add_p(
        doc,
        "Diagrama lógico de la implementación EP1: Internet → ALB (pública) → ASG EC2+Docker "
        "(privada App, Multi-AZ) → EC2 MariaDB (privada Data) + AWS Backup. Detalle Mermaid en "
        "04-architecture-evidence/arquitectura-freshbox.md.",
    )
    add_callout(
        doc,
        "Nota sobre el puerto 443 del enunciado",
        "El enunciado menciona 80/443. En el lab el ALB expone HTTP :80 (sin ACM). La segmentación "
        "ALB→App→Data y el cifrado EBS se mantienen. En la demo se usa siempre http://.",
    )

    add_h2(doc, "1.7.2 TO-BE empresarial sin restricciones de Learner Lab")
    add_p(
        doc,
        "Una vez cumplida la pauta, se recomienda evolucionar la misma topología de tres capas hacia "
        "servicios administrados. Esta sección no reemplaza la demo; reduce riesgo operacional y "
        "carga indiferenciada (parcheo de SO, Docker en hosts, motor DB).",
    )
    add_corp_table(
        doc,
        ["Componente EP1 (lab)", "Evolución TO-BE", "Beneficio"],
        [
            ["EC2 + Docker + ASG", "Amazon ECS con Fargate (5 servicios)", "Escalado por servicio; menos ops de SO"],
            ["MySQL/MariaDB en EC2", "Amazon RDS for MySQL Multi-AZ", "Failover administrado; backups automáticos"],
            ["1 NAT Gateway", "NAT por AZ y/o VPC endpoints (ECR, S3, logs)", "Menor SPOF de salida"],
            ["HTTP :80", "HTTPS (ACM) + AWS WAF delante del ALB", "Cifrado en tránsito y perímetro"],
            ["Secretos en tfvars/env", "AWS Secrets Manager + IAM Roles for Tasks", "Menos filtración de credenciales"],
            ["Validación manual / scripts", "CI/CD con tests, escaneo de imágenes, rollback", "Excelencia operativa"],
            ["Observabilidad básica", "CloudWatch + CloudTrail", "Detección y auditoría"],
            ["VPC /22 (lab)", "Puede ampliarse a /16 en producción", "Espacio de direccionamiento"],
        ],
    )
    add_p(
        doc,
        "Flujo TO-BE de validación funcional: cliente → WAF/ALB → ECS → RDS, con las mismas "
        "operaciones CRUD. En la defensa oral se enfatiza: primero evidencia del lab (ALB→EC2→MySQL); "
        "después se muestra el diagrama TO-BE como hoja de ruta.",
    )

    add_h1(doc, "Conclusiones")
    add_p(
        doc,
        "FreshBox necesita una base cloud alineada a su crecimiento. Este informe valida, en primer "
        "lugar, la implementación EP1 exigida por la pauta: VPC /22, seis subredes, ALB, EC2+Docker "
        "con ASG 2–4, MySQL en EC2 privada, cinco contenedores, SG por capa y CRUD vía ALB, "
        "automatizada con Terraform y GitHub Actions en el Learner Lab.",
    )
    add_p(
        doc,
        "En segundo lugar, propone un TO-BE empresarial (ECS/Fargate, RDS Multi-AZ, WAF, Secrets "
        "Manager, observabilidad y CI/CD) que elimina restricciones del lab y baja el riesgo "
        "operacional. La frase guía de la defensa es: la implementación de laboratorio cumple la "
        "pauta académica; el TO-BE elimina restricciones de Learner Lab y reduce riesgo mediante "
        "servicios administrados.",
    )

    add_h1(doc, "Glosario")
    add_corp_table(
        doc,
        ["Término", "Significado"],
        [
            ["ALB", "Application Load Balancer — balanceador HTTP/HTTPS de aplicación"],
            ["ASG", "Auto Scaling Group — autoescalado de instancias EC2 (pauta EP1)"],
            ["ECS", "Elastic Container Service — orquestación de contenedores (TO-BE)"],
            ["Fargate", "Cómputo serverless para contenedores en ECS (TO-BE)"],
            ["RDS", "Relational Database Service — base administrada Multi-AZ (TO-BE)"],
            ["ECR", "Elastic Container Registry — registro de imágenes Docker"],
            ["VPC", "Virtual Private Cloud — red virtual aislada en AWS"],
            ["SG", "Security Group — firewall stateful por ENI/instancia"],
            ["NAT / IGW", "Salida controlada desde privadas / puerta a Internet"],
            ["CRUD", "Create, Read, Update, Delete"],
            ["IaC", "Infrastructure as Code (Terraform)"],
            ["WAF (AWS Framework)", "Well-Architected Framework — seis pilares de diseño"],
            ["AWS WAF", "Web Application Firewall — protección perimetral (TO-BE)"],
            ["CAPEX / OPEX", "Gasto de capital / gasto operacional"],
            ["Learner Lab", "Entorno AWS Academy con límites (LabRole, regiones, SCP)"],
            ["TO-BE", "Estado objetivo / arquitectura futura"],
            ["AVA / LMS", "Aula digital (concepto) / plataforma (p. ej. Blackboard)"],
        ],
    )

    add_h1(doc, "Referencias")
    refs = [
        "Amazon Web Services. (s. f.). AWS Well-Architected Framework. https://docs.aws.amazon.com/wellarchitected/latest/framework/welcome.html",
        "Amazon Web Services. (s. f.). Amazon EC2 Auto Scaling. https://docs.aws.amazon.com/autoscaling/",
        "Amazon Web Services. (s. f.). Elastic Load Balancing — Application Load Balancers. https://docs.aws.amazon.com/elasticloadbalancing/",
        "Amazon Web Services. (s. f.). Amazon ECS. https://docs.aws.amazon.com/ecs/",
        "Amazon Web Services. (s. f.). Amazon RDS for MySQL. https://docs.aws.amazon.com/AmazonRDS/latest/UserGuide/CHAP_MySQL.html",
        "HashiCorp. (s. f.). Terraform documentation. https://developer.hashicorp.com/terraform/docs",
        "Pastenet, I. A. (2026). Evaluación Parcial N°1 — ARY1102 Arquitectura Cloud (encargo estudiante). Duoc UC.",
        "Matishac14. (2026). ARY1102_Arquitectura_Cloud [Repositorio]. https://github.com/Matishac14/ARY1102_Arquitectura_Cloud",
    ]
    for r in refs:
        p = doc.add_paragraph()
        run = p.add_run(r)
        set_run_font(run, size=10, color=TEXTO)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)

    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2)
        sec.right_margin = Cm(2)

    doc.save(str(OUT))
    print("Wrote", OUT, "size", OUT.stat().st_size)


if __name__ == "__main__":
    main()
